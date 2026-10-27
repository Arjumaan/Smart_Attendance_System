"""Enterprise Attendance Report Generation Service.

Provides:
- CSV report generator (UTF-8 formatted).
- Microsoft Excel (.xlsx) styled report generator via openpyxl.
- Enterprise PDF report generator with corporate header & audit tables via ReportLab.
- High-level query & filter dispatcher.
"""

import csv
from datetime import date, datetime
import io
from typing import List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import or_
from sqlalchemy.orm import Session, joinedload

from backend.app.models.attendance import AttendanceLog
from backend.app.models.department import Department
from backend.app.models.user import User


def _extract_record_fields(log: AttendanceLog) -> dict:
    """Safely extract display fields from an AttendanceLog model."""
    emp_id = "N/A"
    emp_name = "Unknown"
    dept_name = "N/A"

    if log.user:
        emp_id = log.user.employee_id or "N/A"
        emp_name = log.user.full_name or "Unknown"
        if log.user.department:
            dept_name = log.user.department.name or "N/A"

    check_in_str = (
        log.check_in_time.strftime("%Y-%m-%d %H:%M:%S")
        if log.check_in_time
        else "N/A"
    )
    check_out_str = (
        log.check_out_time.strftime("%Y-%m-%d %H:%M:%S")
        if log.check_out_time
        else "N/A"
    )
    hours_str = f"{log.total_hours:.2f}" if log.total_hours is not None else "N/A"
    late_str = "Yes" if log.is_late else "No"

    return {
        "date": str(log.date),
        "employee_id": emp_id,
        "full_name": emp_name,
        "department": dept_name,
        "check_in": check_in_str,
        "check_out": check_out_str,
        "total_hours": hours_str,
        "status": log.status,
        "is_late": late_str,
        "source": log.source or "KIOSK",
    }


def generate_csv_report(records: List[AttendanceLog]) -> bytes:
    """Generate a valid UTF-8 CSV report from attendance records.

    Headers:
    Date, Employee ID, Employee Name, Department, Check-In, Check-Out, Total Hours, Status, Late, Source
    """
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    headers = [
        "Date",
        "Employee ID",
        "Employee Name",
        "Department",
        "Check-In",
        "Check-Out",
        "Total Hours",
        "Status",
        "Late",
        "Source",
    ]
    writer.writerow(headers)

    for rec in records:
        f = _extract_record_fields(rec)
        writer.writerow([
            f["date"],
            f["employee_id"],
            f["full_name"],
            f["department"],
            f["check_in"],
            f["check_out"],
            f["total_hours"],
            f["status"],
            f["is_late"],
            f["source"],
        ])

    return output.getvalue().encode("utf-8")


def generate_excel_report(records: List[AttendanceLog]) -> bytes:
    """Generate a professionally styled Microsoft Excel (.xlsx) workbook using openpyxl.

    Includes:
    - Enterprise navy blue header styling with bold white text.
    - Zebra-striped data rows with thin borders.
    - Auto-adjusted column widths.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"

    headers = [
        "Date",
        "Employee ID",
        "Employee Name",
        "Department",
        "Check-In",
        "Check-Out",
        "Total Hours",
        "Status",
        "Late",
        "Source",
    ]

    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = cell_border

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for row_idx, rec in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        f = _extract_record_fields(rec)
        row_vals = [
            f["date"],
            f["employee_id"],
            f["full_name"],
            f["department"],
            f["check_in"],
            f["check_out"],
            f["total_hours"],
            f["status"],
            f["is_late"],
            f["source"],
        ]
        ws.append(row_vals)

        fill_to_apply = zebra_fill if row_idx % 2 == 0 else white_fill
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill_to_apply
            cell.border = cell_border
            # Column-specific alignment
            if col_idx in (1, 5, 6, 7, 8, 9, 10):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(
    records: List[AttendanceLog],
    start_date: date,
    end_date: date,
    department_name: str = "All",
) -> bytes:
    """Generate an enterprise attendance PDF audit report using ReportLab.

    Includes:
    - Corporate title and report metadata block.
    - Attendance summary KPI statistics.
    - Grid table with colored header, alternating row background, and clean typography.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1E3A8A"),
        fontName="Helvetica-Bold",
        alignment=0,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        fontName="Helvetica",
    )
    meta_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
        fontName="Helvetica",
    )
    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
        fontName="Helvetica",
    )
    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        alignment=1,
    )

    elements = []

    # Title & Header
    elements.append(Paragraph("Enterprise Smart Attendance System", title_style))
    elements.append(Paragraph("Automated Biometrics & Presence Audit Report", subtitle_style))
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1E3A8A"), spaceAfter=12))

    # Metadata & KPI summary
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total_records = len(records)
    late_count = sum(1 for r in records if r.is_late or r.status == "LATE")
    total_work_hours = sum(r.total_hours for r in records if r.total_hours is not None)

    meta_data = [
        [
            Paragraph(f"<b>Date Range:</b> {start_date} to {end_date}", meta_style),
            Paragraph(f"<b>Department:</b> {department_name}", meta_style),
            Paragraph(f"<b>Generated At:</b> {now_str}", meta_style),
        ],
        [
            Paragraph(f"<b>Total Logs:</b> {total_records}", meta_style),
            Paragraph(f"<b>Late Check-Ins:</b> {late_count}", meta_style),
            Paragraph(f"<b>Total Work Hours:</b> {total_work_hours:.2f} hrs", meta_style),
        ],
    ]
    meta_table = Table(meta_data, colWidths=[240, 240, 240])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 14))

    # Main Attendance Records Table
    col_headers = [
        "Date",
        "Emp ID",
        "Employee Name",
        "Department",
        "Check-In",
        "Check-Out",
        "Hours",
        "Status",
        "Late",
        "Source",
    ]
    table_rows = [[Paragraph(h, table_header_style) for h in col_headers]]

    for rec in records:
        f = _extract_record_fields(rec)
        table_rows.append([
            Paragraph(f["date"], table_cell_style),
            Paragraph(f["employee_id"], table_cell_style),
            Paragraph(f["full_name"], table_cell_style),
            Paragraph(f["department"], table_cell_style),
            Paragraph(f["check_in"], table_cell_style),
            Paragraph(f["check_out"], table_cell_style),
            Paragraph(f["total_hours"], table_cell_style),
            Paragraph(f["status"], table_cell_style),
            Paragraph(f["is_late"], table_cell_style),
            Paragraph(f["source"], table_cell_style),
        ])

    if total_records == 0:
        no_data_para = Paragraph("No attendance records found for the selected criteria.", meta_style)
        table_rows.append([no_data_para] + [Paragraph("", table_cell_style) for _ in range(len(col_headers) - 1)])

    # 10 columns across ~720 points printable width in landscape letter
    col_widths = [62, 55, 95, 80, 95, 95, 45, 52, 38, 50]
    rec_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
    rec_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(rec_table)

    doc.build(elements)
    return buf.getvalue()


def generate_report(
    db: Session,
    start_date: date,
    end_date: date,
    department_id: Optional[int] = None,
    employee_id: Optional[str] = None,
    status: Optional[str] = None,
    report_format: str = "csv",
) -> Tuple[bytes, str, str]:
    """Query database with filters and export report in CSV, XLSX, or PDF.

    Returns:
        Tuple of (content_bytes, media_type, filename)
    """
    query = (
        db.query(AttendanceLog)
        .join(User, AttendanceLog.user_id == User.id)
        .options(joinedload(AttendanceLog.user).joinedload(User.department))
        .filter(AttendanceLog.date >= start_date, AttendanceLog.date <= end_date)
    )

    dept_name = "All Departments"
    if department_id:
        query = query.filter(User.department_id == department_id)
        dept_obj = db.query(Department).filter(Department.id == department_id).first()
        if dept_obj:
            dept_name = dept_obj.name

    if employee_id and employee_id.strip():
        term = f"%{employee_id.strip()}%"
        query = query.filter(
            or_(
                User.employee_id.ilike(term),
                User.full_name.ilike(term),
                User.email.ilike(term),
            )
        )

    records = None
    if status and status.strip():
        stat_upper = status.strip().upper()
        if stat_upper == "LATE":
            query = query.filter(
                or_(
                    AttendanceLog.status == "LATE",
                    AttendanceLog.is_late.is_(True),
                )
            )
        elif stat_upper == "PRESENT":
            query = query.filter(
                AttendanceLog.status == "PRESENT",
                AttendanceLog.is_late.is_(False),
            )
        elif stat_upper == "ABSENT":
            # For absentees, generate synthetic attendance records for unlogged active employees
            target_date = end_date or start_date or date.today()
            active_q = (
                db.query(User)
                .options(joinedload(User.department))
                .filter(User.is_active.is_(True))
            )
            if department_id:
                active_q = active_q.filter(User.department_id == department_id)
            if employee_id and employee_id.strip():
                clean_term = f"%{employee_id.strip()}%"
                active_q = active_q.filter(
                    or_(
                        User.employee_id.ilike(clean_term),
                        User.full_name.ilike(clean_term),
                        User.email.ilike(clean_term),
                    )
                )

            logged_uids = {
                uid
                for (uid,) in db.query(AttendanceLog.user_id)
                .filter(AttendanceLog.date == target_date)
                .all()
            }

            absent_users = [u for u in active_q.all() if u.id not in logged_uids]
            now_dt = datetime.now()
            records = [
                AttendanceLog(
                    id=-u.id,
                    user_id=u.id,
                    user=u,
                    date=target_date,
                    check_in_time=None,
                    check_out_time=None,
                    total_hours=0.0,
                    status="ABSENT",
                    is_late=False,
                    source="SYSTEM",
                    notes="No presence recorded.",
                    created_at=now_dt,
                    updated_at=now_dt,
                )
                for u in absent_users
            ]
        else:
            query = query.filter(AttendanceLog.status == stat_upper)

    if records is None:
        records = (
            query.order_by(AttendanceLog.date.desc(), AttendanceLog.check_in_time.asc())
            .all()
        )

    fmt = report_format.lower().strip()
    date_str = f"{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}"

    if fmt in ("xlsx", "excel"):
        content = generate_excel_report(records)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"attendance_report_{date_str}.xlsx"
    elif fmt == "pdf":
        content = generate_pdf_report(records, start_date, end_date, department_name=dept_name)
        media_type = "application/pdf"
        filename = f"attendance_report_{date_str}.pdf"
    else:
        content = generate_csv_report(records)
        media_type = "text/csv; charset=utf-8"
        filename = f"attendance_report_{date_str}.csv"

    return content, media_type, filename
